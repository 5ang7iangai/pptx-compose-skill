"""
text/material.py — 多格式材料统一解析

支持格式：
  - 纯文字 (.txt / 直传字符串)
  - Word (.docx)
  - PDF (.pdf)
  - XMind (.xmind)
  - PPT (.pptx)

统一输出 JSON 结构：
{
  "source_type": "docx",
  "source_path": "材料.docx",
  "title": "提取到的标题",
  "sections": [
    {"heading": "一、背景", "body": "...", "level": 1}
  ],
  "images": ["desc: 架构图（第3页）"],
  "tables": [{"caption": "", "headers": [...], "rows": [[...]]}],
  "raw_char_count": 12345
}
"""

from __future__ import annotations

import json
import logging
import re
import zipfile
from pathlib import Path
from typing import Any

log = logging.getLogger(__name__)


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------

def read_material(path: str | Path | None = None, *, text: str | None = None) -> dict:
    """
    解析材料文件，返回统一结构的 dict。
    传 path 解析文件，传 text 直接当纯文字处理。
    """
    if text is not None:
        return _parse_plain_text(text, source="(直传文本)")

    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"材料文件不存在: {path}")

    suffix = path.suffix.lower()
    log.info("解析材料: %s (格式: %s)", path.name, suffix)

    if suffix == ".json":
        data = json.loads(path.read_text(encoding="utf-8"))
        log.info("加载已解析的材料 JSON: %s", path.name)
        return data
    elif suffix == ".txt":
        return _parse_plain_text(path.read_text(encoding="utf-8"), source=str(path))
    elif suffix == ".docx":
        return _parse_docx(path)
    elif suffix == ".pdf":
        return _parse_pdf(path)
    elif suffix == ".xmind":
        return _parse_xmind(path)
    elif suffix == ".pptx":
        return _parse_pptx(path)
    else:
        log.warning("未知格式 %s，当纯文本处理", suffix)
        return _parse_plain_text(path.read_text(encoding="utf-8"), source=str(path))


def save_material(material: dict, output_path: str | Path) -> Path:
    """把解析结果保存为 JSON"""
    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(material, ensure_ascii=False, indent=2), encoding="utf-8")
    log.info("材料解析结果已保存: %s (%d sections, %d chars)",
             out, len(material.get("sections", [])), material.get("raw_char_count", 0))
    return out


# ---------------------------------------------------------------------------
# Plain text
# ---------------------------------------------------------------------------

def _parse_plain_text(text: str, source: str = "") -> dict:
    sections = _split_text_to_sections(text)
    return {
        "source_type": "text",
        "source_path": source,
        "title": _guess_title(text),
        "sections": sections,
        "images": [],
        "tables": [],
        "raw_char_count": len(text),
    }


def _split_text_to_sections(text: str) -> list[dict]:
    """按常见中文标题模式分段"""
    lines = text.splitlines()
    sections: list[dict] = []
    current_heading = ""
    current_level = 0
    current_body_lines: list[str] = []

    heading_pattern = re.compile(
        r'^(?:'
        r'[一二三四五六七八九十]+[、．.]'   # 一、 二、
        r'|(?:\d+\.)+\s*'                    # 1. 1.1 1.1.1
        r'|[（(]\s*[一二三四五六七八九十\d]+\s*[)）]'  # （一） (1)
        r'|第[一二三四五六七八九十\d]+[章节部分篇]'    # 第一章
        r'|#{1,6}\s+'                         # Markdown heading
        r')'
    )

    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue
        if heading_pattern.match(stripped) and len(stripped) < 80:
            if current_heading or current_body_lines:
                sections.append({
                    "heading": current_heading,
                    "body": "\n".join(current_body_lines).strip(),
                    "level": current_level,
                })
            current_heading = stripped
            current_level = _guess_heading_level(stripped)
            current_body_lines = []
        else:
            current_body_lines.append(stripped)

    if current_heading or current_body_lines:
        sections.append({
            "heading": current_heading,
            "body": "\n".join(current_body_lines).strip(),
            "level": current_level,
        })

    if not sections and text.strip():
        sections.append({"heading": "", "body": text.strip(), "level": 0})

    return sections


def _guess_heading_level(heading: str) -> int:
    if re.match(r'^#{1,6}', heading):
        return heading.index(' ') if ' ' in heading else len(heading.split()[0])
    if re.match(r'^[一二三四五六七八九十]+[、．.]', heading):
        return 1
    if re.match(r'^[（(][一二三四五六七八九十]+[)）]', heading):
        return 2
    if re.match(r'^\d+\.\d+\.\d+', heading):
        return 3
    if re.match(r'^\d+\.\d+', heading):
        return 2
    if re.match(r'^\d+\.', heading):
        return 1
    return 1


def _guess_title(text: str) -> str:
    for line in text.splitlines()[:10]:
        stripped = line.strip()
        if stripped and len(stripped) < 60:
            return stripped
    return ""


# ---------------------------------------------------------------------------
# Word (.docx)
# ---------------------------------------------------------------------------

def _parse_docx(path: Path) -> dict:
    try:
        from docx import Document
    except ImportError:
        raise ImportError("解析 .docx 需要 python-docx：pip install python-docx")

    doc = Document(str(path))
    sections: list[dict] = []
    images: list[str] = []
    tables: list[dict] = []
    full_text_parts: list[str] = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        full_text_parts.append(text)

        style_name = (para.style.name or "").lower()
        if "heading" in style_name:
            level = int(re.search(r'\d+', style_name).group()) if re.search(r'\d+', style_name) else 1
            sections.append({"heading": text, "body": "", "level": level})
        elif sections:
            sections[-1]["body"] += ("\n" if sections[-1]["body"] else "") + text
        else:
            sections.append({"heading": "", "body": text, "level": 0})

    for i, table in enumerate(doc.tables):
        headers = [cell.text.strip() for cell in table.rows[0].cells] if table.rows else []
        rows = []
        for row in table.rows[1:]:
            rows.append([cell.text.strip() for cell in row.cells])
        tables.append({"caption": f"表{i+1}", "headers": headers, "rows": rows})

    for rel in doc.part.rels.values():
        if "image" in rel.reltype:
            images.append(f"desc: 嵌入图片 ({rel.target_ref})")

    full_text = "\n".join(full_text_parts)
    title = doc.core_properties.title or _guess_title(full_text)

    return {
        "source_type": "docx",
        "source_path": str(path),
        "title": title or path.stem,
        "sections": sections,
        "images": images,
        "tables": tables,
        "raw_char_count": len(full_text),
    }


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

def _parse_pdf(path: Path) -> dict:
    try:
        import fitz  # PyMuPDF
    except ImportError:
        raise ImportError("解析 .pdf 需要 PyMuPDF：pip install PyMuPDF")

    doc = fitz.open(str(path))
    full_text_parts: list[str] = []
    images: list[str] = []

    for page_num in range(len(doc)):
        page = doc[page_num]
        text = page.get_text("text").strip()
        if text:
            full_text_parts.append(text)
        img_list = page.get_images(full=True)
        for img in img_list:
            images.append(f"desc: 图片 (第{page_num+1}页, xref={img[0]})")

    doc.close()
    full_text = "\n\n".join(full_text_parts)
    sections = _split_text_to_sections(full_text)

    return {
        "source_type": "pdf",
        "source_path": str(path),
        "title": _guess_title(full_text) or path.stem,
        "sections": sections,
        "images": images,
        "tables": [],
        "raw_char_count": len(full_text),
    }


# ---------------------------------------------------------------------------
# XMind (.xmind)
# ---------------------------------------------------------------------------

def _parse_xmind(path: Path) -> dict:
    """XMind 文件是 ZIP，内含 content.json（新版）或 content.xml（旧版）"""
    with zipfile.ZipFile(str(path), "r") as zf:
        names = zf.namelist()

        if "content.json" in names:
            raw = json.loads(zf.read("content.json").decode("utf-8"))
            sections = _xmind_json_to_sections(raw)
            full_text = "\n".join(s["heading"] + " " + s["body"] for s in sections)
        elif any(n.endswith("content.xml") for n in names):
            xml_name = next(n for n in names if n.endswith("content.xml"))
            from xml.etree import ElementTree as ET
            tree = ET.fromstring(zf.read(xml_name))
            sections = _xmind_xml_to_sections(tree)
            full_text = "\n".join(s["heading"] + " " + s["body"] for s in sections)
        else:
            log.warning("XMind 文件中找不到 content.json 或 content.xml: %s", names[:10])
            return _parse_plain_text("", source=str(path))

    return {
        "source_type": "xmind",
        "source_path": str(path),
        "title": sections[0]["heading"] if sections else path.stem,
        "sections": sections,
        "images": [],
        "tables": [],
        "raw_char_count": len(full_text),
    }


def _xmind_json_to_sections(data: Any) -> list[dict]:
    """解析 XMind 新版 content.json 的层级结构"""
    sections: list[dict] = []

    def walk(node: dict, level: int = 0):
        title = node.get("title", "").strip()
        if not title:
            return
        children = node.get("children", {}).get("attached", [])
        child_texts = [c.get("title", "").strip() for c in children if c.get("title", "").strip()]
        sections.append({
            "heading": title,
            "body": "；".join(child_texts) if child_texts else "",
            "level": level,
        })
        for child in children:
            if child.get("children"):
                walk(child, level + 1)

    if isinstance(data, list):
        for sheet in data:
            root = sheet.get("rootTopic", {})
            walk(root, 0)
    elif isinstance(data, dict):
        root = data.get("rootTopic", {})
        walk(root, 0)

    return sections


def _xmind_xml_to_sections(root) -> list[dict]:
    """解析 XMind 旧版 content.xml 的层级结构"""
    ns = {"xm": "urn:xmind:xmap:xmlns:content:2.0"}
    sections: list[dict] = []

    def walk(topic, level=0):
        title_el = topic.find("xm:title", ns)
        title = title_el.text.strip() if title_el is not None and title_el.text else ""
        if not title:
            return
        child_texts = []
        for ch in topic.findall("xm:children/xm:topics/xm:topic", ns):
            t = ch.find("xm:title", ns)
            if t is not None and t.text:
                child_texts.append(t.text.strip())
        sections.append({
            "heading": title,
            "body": "；".join(child_texts),
            "level": level,
        })
        for ch in topic.findall("xm:children/xm:topics/xm:topic", ns):
            walk(ch, level + 1)

    for sheet in root.findall(".//xm:sheet", ns) or [root]:
        for topic in sheet.findall(".//xm:topic", ns)[:1]:
            walk(topic, 0)

    return sections


# ---------------------------------------------------------------------------
# PPT (.pptx)
# ---------------------------------------------------------------------------

def _parse_pptx(path: Path) -> dict:
    try:
        from pptx import Presentation
    except ImportError:
        raise ImportError("解析 .pptx 需要 python-pptx：pip install python-pptx")

    prs = Presentation(str(path))
    sections: list[dict] = []
    images: list[str] = []
    full_text_parts: list[str] = []

    for i, slide in enumerate(prs.slides, 1):
        slide_texts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = para.text.strip()
                    if text:
                        slide_texts.append(text)
            if shape.shape_type == 13:  # Picture
                images.append(f"desc: 图片 (第{i}页)")

        if slide_texts:
            heading = slide_texts[0] if len(slide_texts[0]) < 60 else ""
            body = "\n".join(slide_texts[1:] if heading else slide_texts)
            sections.append({
                "heading": f"第{i}页: {heading}".strip(": "),
                "body": body,
                "level": 1,
            })
            full_text_parts.extend(slide_texts)

        notes = slide.notes_slide.notes_text_frame.text.strip() if slide.has_notes_slide else ""
        if notes:
            sections[-1]["body"] += f"\n[备注] {notes}" if sections else ""

    full_text = "\n".join(full_text_parts)
    return {
        "source_type": "pptx",
        "source_path": str(path),
        "title": _guess_title(full_text) or path.stem,
        "sections": sections,
        "images": images,
        "tables": [],
        "raw_char_count": len(full_text),
    }


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

if __name__ == "__main__":
    import sys

    logging.basicConfig(level=logging.INFO)

    if len(sys.argv) < 2:
        print("用法: python -m text.material <材料文件> [输出.json]")
        sys.exit(1)

    src = Path(sys.argv[1])
    out = Path(sys.argv[2]) if len(sys.argv) > 2 else src.with_suffix(".material.json")

    result = read_material(src)
    save_material(result, out)
    print(f"[OK] 解析完成: {len(result['sections'])} sections, "
          f"{result['raw_char_count']} chars, "
          f"{len(result['images'])} images, "
          f"{len(result['tables'])} tables")
